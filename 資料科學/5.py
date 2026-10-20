
from selenium import webdriver
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import re
import time
import os

# 台北週末票房排行榜
START_URL = "https://app2.atmovies.com.tw/boxoffice/TWweekend/"
OUTPUT_FILE = "Taipei_movies.xlsx"


def clean_money(text):
    """將票房文字轉成數字，例如 $4,019 -> 4019"""
    text = text.replace(",", "")
    match = re.search(r"\$?\s*(\d+(?:\.\d+)?)", text)

    if match:
        return float(match.group(1))

    return None


def main():
    driver = webdriver.Chrome()
    movies = []

    try:
        print("正在開啟電影票房網站...")
        driver.get(START_URL)
        time.sleep(3)

        soup = BeautifulSoup(driver.page_source, "html.parser")

        # 尋找通往完整榜單的 more 網址
        more_url = None

        for a in soup.find_all("a", href=True):
            href = a["href"]

            if "/boxoffice/" in href and "/more/" in href:
                more_url = href
                break

        if more_url:
            if more_url.startswith("/"):
                more_url = "https://app2.atmovies.com.tw" + more_url
            elif not more_url.startswith("http"):
                more_url = "https://app2.atmovies.com.tw/" + more_url

            driver.get(more_url)
            time.sleep(3)

        soup = BeautifulSoup(driver.page_source, "html.parser")

        # 依序分析表格的每一列
        rows = soup.find_all("tr")

        for i, row in enumerate(rows):
            cells = row.find_all(["td", "th"], recursive=False)
            texts = [
                cell.get_text(" ", strip=True)
                for cell in cells
            ]

            if not texts:
                continue

            # 找出以排名開頭、且有電影名稱的列
            rank_text = texts[0].strip()

            if not re.fullmatch(r"\d+", rank_text):
                continue

            rank = int(rank_text)

            if rank < 1 or rank > 20:
                continue

            title = ""
            for cell in cells[1:]:
                title = cell.get_text(" ", strip=True)
                if title:
                    break

            if not title:
                continue

            # 電影名稱列與票房列可能分開，
            # 所以往後尋找下一個排名列之前的票房數字
            money_values = []

            for next_row in rows[i + 1:]:
                next_cells = next_row.find_all(
                    ["td", "th"], recursive=False
                )
                next_texts = [
                    c.get_text(" ", strip=True)
                    for c in next_cells
                ]

                if not next_texts:
                    continue

                # 遇到下一部電影就停止
                if re.fullmatch(r"\d+", next_texts[0].strip()):
                    break

                row_text = " ".join(next_texts)

                # 只擷取有 $ 符號的票房數字
                amounts = re.findall(
                    r"\$\s*[\d,]+(?:\.\d+)?", row_text
                )

                for amount in amounts:
                    value = clean_money(amount)
                    if value is not None:
                        money_values.append(value)

            if len(money_values) >= 2:
                weekly = money_values[0]
                cumulative = money_values[1]

                movies.append([
                    rank, title, weekly, cumulative
                ])

                print(
                    f"排名：{rank}｜{title}｜"
                    f"本週票房：{weekly:,.0f} 萬｜"
                    f"累計票房：{cumulative:,.0f} 萬"
                )

        # 依排名排序並去除重複排名
        unique_movies = {}
        for movie in movies:
            unique_movies[movie[0]] = movie

        movies = [
            unique_movies[rank]
            for rank in sorted(unique_movies)
        ][:20]

        if not movies:
            print("沒有抓到電影資料，請確認網站結構是否有變更。")
            return

        # 建立 Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "台北週末票房"

        headers = [
            "排名",
            "電影名稱",
            "本週票房（萬新台幣）",
            "累計票房（萬新台幣）"
        ]
        ws.append(headers)

        for movie in movies:
            ws.append(movie)

        # 設定標題格式
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="4472C4"
            )
            cell.alignment = Alignment(horizontal="center")

        # 調整欄寬
        widths = [10, 65, 25, 25]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[
                get_column_letter(index)
            ].width = width

        # 票房顯示千分位
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
            for cell in row:
                cell.number_format = '#,##0'

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # 存到程式所在資料夾
        output_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            OUTPUT_FILE
        )
        wb.save(output_path)

        print("\n爬蟲完成！")
        print(f"成功擷取 {len(movies)} 部電影")
        print(f"Excel 儲存位置：{output_path}")

    finally:
        driver.quit()


if __name__ == "__main__":
    main()
